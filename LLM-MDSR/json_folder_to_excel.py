#!/usr/bin/env python3
"""Batch-convert JSON files in a folder into one Excel table."""

from __future__ import annotations

import argparse
import ast
import json
import re
from pathlib import Path
from typing import Any

import pandas as pd
import sympy as sp


PREFERRED_COLUMNS = [
    "source_file",
    "source_index",
    "sample_order",
    "function",
    "formula",
    "complexity",
    "formula_error",
    "score",
    "mean_nmse",
    "min_nmse",
    "max_nmse",
    "median_nmse",
    "std_nmse",
    "worst2_mean_nmse",
    "sample_time",
    "evaluate_time",
]


class SymbolicFunctionParser:
    """Safely convert a restricted NumPy-style function body to SymPy."""

    FUNCTIONS = {
        "sin": sp.sin, "cos": sp.cos, "tan": sp.tan,
        "arcsin": sp.asin, "asin": sp.asin,
        "arccos": sp.acos, "acos": sp.acos,
        "arctan": sp.atan, "atan": sp.atan,
        "sinh": sp.sinh, "cosh": sp.cosh, "tanh": sp.tanh,
        "exp": sp.exp, "log": sp.log, "sqrt": sp.sqrt,
        "abs": sp.Abs, "fabs": sp.Abs, "sign": sp.sign,
        "maximum": sp.Max, "minimum": sp.Min,
    }

    def __init__(self) -> None:
        self.env: dict[str, sp.Expr] = {}

    @staticmethod
    def dotted_name(node: ast.AST) -> str:
        if isinstance(node, ast.Name):
            return node.id
        if isinstance(node, ast.Attribute):
            left = SymbolicFunctionParser.dotted_name(node.value)
            return f"{left}.{node.attr}" if left else node.attr
        return ""

    def expression(self, node: ast.AST) -> sp.Expr:
        if isinstance(node, ast.Constant):
            if isinstance(node.value, int):
                return sp.Integer(node.value)
            if isinstance(node.value, float):
                # Normalize 1.0 -> 1 so exact symbolic simplification is not
                # blocked by an otherwise unnecessary floating-point atom.
                if node.value.is_integer():
                    return sp.Integer(int(node.value))
                return sp.Float(node.value)
            raise ValueError(f"不支持的常量：{node.value!r}")

        if isinstance(node, ast.Name):
            if node.id == "pi":
                return sp.pi
            return self.env.get(node.id, sp.Symbol(node.id, real=True))

        if isinstance(node, ast.Attribute):
            name = self.dotted_name(node)
            if name in {"np.pi", "numpy.pi"}:
                return sp.pi
            if name in {"np.e", "numpy.e"}:
                return sp.E
            raise ValueError(f"不支持的属性：{name}")

        if isinstance(node, ast.BinOp):
            left, right = self.expression(node.left), self.expression(node.right)
            operations = {
                ast.Add: lambda: left + right,
                ast.Sub: lambda: left - right,
                ast.Mult: lambda: left * right,
                ast.Div: lambda: left / right,
                ast.Pow: lambda: left ** right,
                ast.Mod: lambda: sp.Mod(left, right),
            }
            operation = operations.get(type(node.op))
            if operation is None:
                raise ValueError(f"不支持的二元运算：{type(node.op).__name__}")
            return operation()

        if isinstance(node, ast.UnaryOp):
            value = self.expression(node.operand)
            if isinstance(node.op, ast.USub):
                return -value
            if isinstance(node.op, ast.UAdd):
                return value
            if isinstance(node.op, ast.Not):
                return sp.Not(value)
            raise ValueError(f"不支持的一元运算：{type(node.op).__name__}")

        if isinstance(node, ast.Subscript):
            base = self.dotted_name(node.value)
            index_node = node.slice
            if isinstance(index_node, ast.Constant) and isinstance(index_node.value, int):
                index = index_node.value
            else:
                index = self.expression(index_node)
            # Dataset parameters get compact mathematical names p0, p1, ...
            if base == "params" and isinstance(index, int):
                return sp.Symbol(f"p{index}", real=True)
            return sp.Symbol(f"{base}[{index}]", real=True)

        if isinstance(node, ast.Call):
            full_name = self.dotted_name(node.func)
            short_name = full_name.rsplit(".", 1)[-1]
            args = [self.expression(arg) for arg in node.args]

            if short_name in self.FUNCTIONS:
                return self.FUNCTIONS[short_name](*args)
            if short_name == "square" and len(args) == 1:
                return args[0] ** 2
            if short_name in {"pow", "power"} and len(args) == 2:
                return args[0] ** args[1]
            if short_name == "where" and len(args) == 3:
                return sp.Piecewise((args[1], args[0]), (args[2], True))
            if short_name == "clip" and len(args) == 3:
                return sp.Min(sp.Max(args[0], args[1]), args[2])
            raise ValueError(f"不支持的函数调用：{full_name}")

        if isinstance(node, ast.Compare) and len(node.ops) == len(node.comparators) == 1:
            left, right = self.expression(node.left), self.expression(node.comparators[0])
            relations = {
                ast.Lt: sp.Lt, ast.LtE: sp.Le, ast.Gt: sp.Gt,
                ast.GtE: sp.Ge, ast.Eq: sp.Eq, ast.NotEq: sp.Ne,
            }
            relation = relations.get(type(node.ops[0]))
            if relation is None:
                raise ValueError("不支持的比较运算")
            return relation(left, right)

        raise ValueError(f"不支持的表达式节点：{type(node).__name__}")

    def assign(self, target: ast.AST, value_node: ast.AST) -> None:
        if isinstance(target, (ast.Tuple, ast.List)):
            if not isinstance(value_node, (ast.Tuple, ast.List)):
                raise ValueError("多变量赋值的左右数量无法对应")
            if len(target.elts) != len(value_node.elts):
                raise ValueError("多变量赋值的左右数量不一致")
            for child_target, child_value in zip(target.elts, value_node.elts):
                self.assign(child_target, child_value)
            return
        if not isinstance(target, ast.Name):
            raise ValueError(f"不支持的赋值目标：{type(target).__name__}")
        self.env[target.id] = self.expression(value_node)

    def parse(self, source: str) -> sp.Expr:
        tree = ast.parse(source)
        functions = [node for node in tree.body if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))]
        if not functions:
            raise ValueError("没有找到函数定义")

        for statement in functions[0].body:
            if isinstance(statement, ast.Assign):
                for target in statement.targets:
                    if isinstance(target, ast.Name):
                        self.env[target.id] = self.expression(statement.value)
                    else:
                        self.assign(target, statement.value)
            elif isinstance(statement, ast.AnnAssign) and statement.value is not None:
                self.assign(statement.target, statement.value)
            elif isinstance(statement, ast.AugAssign) and isinstance(statement.target, ast.Name):
                old = self.env.get(statement.target.id, sp.Symbol(statement.target.id, real=True))
                right = self.expression(statement.value)
                if isinstance(statement.op, ast.Add):
                    self.env[statement.target.id] = old + right
                elif isinstance(statement.op, ast.Sub):
                    self.env[statement.target.id] = old - right
                elif isinstance(statement.op, ast.Mult):
                    self.env[statement.target.id] = old * right
                elif isinstance(statement.op, ast.Div):
                    self.env[statement.target.id] = old / right
                else:
                    raise ValueError(f"不支持的复合赋值：{type(statement.op).__name__}")
            elif isinstance(statement, ast.Return):
                if statement.value is None:
                    raise ValueError("return 没有返回表达式")
                return sp.simplify(self.expression(statement.value))
            elif isinstance(statement, (ast.Expr, ast.Pass)):
                # Ignore docstrings and other non-assignment expression lines.
                continue
            else:
                raise ValueError(f"不支持的函数体语句：{type(statement).__name__}")

        raise ValueError("函数中没有找到 return 语句")


def function_to_formula_and_complexity(source: str) -> tuple[str, int]:
    """Return the simplified formula and its SymPy operation count."""
    expression = SymbolicFunctionParser().parse(source)
    formula = str(expression)
    complexity = int(sp.count_ops(expression, visual=False))
    return formula, complexity


def function_to_formula(source: str) -> str:
    """Resolve intermediate assignments and return a simplified formula."""
    formula, _ = function_to_formula_and_complexity(source)
    return formula


def natural_key(path: Path) -> list[Any]:
    """Sort samples_2.json before samples_10.json."""
    return [int(part) if part.isdigit() else part.lower()
            for part in re.split(r"(\d+)", path.name)]


def flatten(value: Any, prefix: str = "") -> dict[str, Any]:
    """Flatten nested dictionaries; keep lists as readable JSON text."""
    if isinstance(value, dict):
        result: dict[str, Any] = {}
        for key, child in value.items():
            column = f"{prefix}.{key}" if prefix else str(key)
            result.update(flatten(child, column))
        return result

    if isinstance(value, list):
        return {prefix or "value": json.dumps(value, ensure_ascii=False)}

    return {prefix or "value": value}


def read_records(path: Path, root: Path) -> list[dict[str, Any]]:
    """Read one JSON file and return one or more flattened table rows."""
    with path.open("r", encoding="utf-8-sig") as file:
        data = json.load(file)

    relative_name = str(path.relative_to(root))
    items = data if isinstance(data, list) else [data]
    rows: list[dict[str, Any]] = []

    for index, item in enumerate(items):
        row = flatten(item)
        row["source_file"] = relative_name
        if isinstance(data, list):
            row["source_index"] = index
        rows.append(row)

    return rows


def reorder_columns(frame: pd.DataFrame) -> pd.DataFrame:
    first = [column for column in PREFERRED_COLUMNS if column in frame.columns]
    remaining = [column for column in frame.columns if column not in first]
    return frame[first + remaining]


def show_progress(completed: int, total: int, message: str = "") -> None:
    """Display a dependency-free, single-line terminal progress bar."""
    width = 30
    ratio = completed / total if total else 1.0
    filled = min(width, int(width * ratio))
    bar = "█" * filled + "-" * (width - filled)
    percent = ratio * 100
    # Limit/pad the status text so a shorter filename overwrites the old one.
    status = message.replace("\n", " ")
    if len(status) > 48:
        status = "..." + status[-45:]
    print(
        f"\r处理进度 [{bar}] {completed:>{len(str(total))}}/{total} "
        f"({percent:6.2f}%)  {status:<48}",
        end="",
        flush=True,
    )


def convert_folder(
    input_dir: Path,
    output_file: Path,
    recursive: bool = False,
    top_n: int = 30,
) -> tuple[int, int, int]:
    if top_n <= 0:
        raise ValueError("top_n 必须是大于 0 的整数。")

    pattern = "**/*.json" if recursive else "*.json"
    json_files = sorted(input_dir.glob(pattern), key=natural_key)
    if not json_files:
        raise FileNotFoundError(f"在文件夹中没有找到 JSON 文件：{input_dir}")

    rows: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []

    total_files = len(json_files)
    show_progress(0, total_files, "准备开始")

    for file_number, path in enumerate(json_files, start=1):
        show_progress(file_number - 1, total_files, f"正在处理：{path.name}")
        try:
            rows.extend(read_records(path, input_dir))
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            errors.append({"source_file": str(path.relative_to(input_dir)),
                           "error": str(exc)})
        show_progress(file_number, total_files, f"已处理：{path.name}")

    print()  # End the progress-bar line.

    if not rows:
        raise ValueError("没有成功读取任何 JSON 记录。")

    all_rows_count = len(rows)
    frame = pd.DataFrame(rows)
    if "score" not in frame.columns:
        raise KeyError("JSON 记录中没有找到用于排名的 score 字段。")

    # Preserve the original score column, but use a numeric helper for robust
    # descending ranking. Stable sorting keeps source-file order for ties.
    frame["_score_numeric"] = pd.to_numeric(frame["score"], errors="coerce")
    valid_score_count = int(frame["_score_numeric"].notna().sum())
    if valid_score_count == 0:
        raise ValueError("score 字段中没有可用于排名的数值。")

    selected_count = min(top_n, valid_score_count)
    print(
        f"正在按 score 从高到低筛选前 {selected_count} 条"
        f"（候选记录共 {all_rows_count} 条）……",
        flush=True,
    )
    frame = (
        frame.loc[frame["_score_numeric"].notna()]
        .sort_values(
            "_score_numeric", ascending=False, kind="stable", na_position="last"
        )
        .head(top_n)
        .drop(columns="_score_numeric")
        .reset_index(drop=True)
    )

    if "function" not in frame.columns:
        raise KeyError("JSON 记录中没有找到 function 字段。")

    # SymPy is deliberately run only after Top-N selection because symbolic
    # simplification is much more expensive than loading and ranking JSON.
    frame["formula"] = None
    frame["complexity"] = None
    formula_errors: list[str | None] = [None] * len(frame)
    print(f"正在转换并化简入选的 {len(frame)} 个公式：", flush=True)
    show_progress(0, len(frame), "准备开始")
    for row_number, function_source in enumerate(frame["function"], start=1):
        source_name = str(frame.at[row_number - 1, "source_file"])
        show_progress(row_number - 1, len(frame), f"正在转换：{source_name}")
        if isinstance(function_source, str) and function_source.strip():
            try:
                formula, complexity = function_to_formula_and_complexity(function_source)
                frame.at[row_number - 1, "formula"] = formula
                frame.at[row_number - 1, "complexity"] = complexity
            except (SyntaxError, TypeError, ValueError, NotImplementedError) as exc:
                formula_errors[row_number - 1] = str(exc)
        else:
            formula_errors[row_number - 1] = "function 字段为空或不是字符串"
        show_progress(row_number, len(frame), f"已转换：{source_name}")
    print()

    if any(error is not None for error in formula_errors):
        frame["formula_error"] = formula_errors

    frame = reorder_columns(frame)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    print("正在写入 Excel 文件……", flush=True)
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="JSON_Data", index=False)
        sheet = writer.sheets["JSON_Data"]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        # Sensible widths; long function/code cells wrap instead of making a
        # sheet thousands of pixels wide.
        for cells in sheet.columns:
            letter = cells[0].column_letter
            header = str(cells[0].value or "")
            if header == "function":
                width = 70
            else:
                width = min(max(len(str(cell.value or "")) for cell in cells) + 2, 35)
            sheet.column_dimensions[letter].width = max(width, 12)

        from openpyxl.styles import Alignment, Font, PatternFill

        header_fill = PatternFill("solid", fgColor="4472C4")
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        if errors:
            pd.DataFrame(errors).to_excel(writer, sheet_name="Errors", index=False)
            error_sheet = writer.sheets["Errors"]
            error_sheet.freeze_panes = "A2"
            error_sheet.column_dimensions["A"].width = 40
            error_sheet.column_dimensions["B"].width = 90

    return len(json_files), all_rows_count, len(frame)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="把指定文件夹中的 JSON 文件汇总为一个 Excel 表格。"
    )
    parser.add_argument("input_dir", type=Path, help="包含 JSON 文件的文件夹")
    parser.add_argument(
        "-o", "--output", type=Path,
        default=Path(__file__).resolve().parent / "json_summary.xlsx",
        help="输出 Excel 路径（默认：脚本目录中的 json_summary.xlsx）",
    )
    parser.add_argument(
        "-r", "--recursive", action="store_true",
        help="递归读取所有子文件夹中的 JSON 文件",
    )
    parser.add_argument(
        "-n", "--top-n", type=int, default=30,
        help="仅保留 score 最高的前 N 条记录（默认：30）",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    input_dir = args.input_dir.expanduser().resolve()
    output_file = args.output.expanduser().resolve()

    if not input_dir.is_dir():
        raise NotADirectoryError(f"输入路径不是文件夹：{input_dir}")

    file_count, candidate_count, selected_count = convert_folder(
        input_dir=input_dir,
        output_file=output_file,
        recursive=args.recursive,
        top_n=args.top_n,
    )
    print(
        f"完成：读取 {file_count} 个 JSON 文件、{candidate_count} 条候选记录，"
        f"写入 score 最高的 {selected_count} 行。"
    )
    print(f"输出文件：{output_file}")


if __name__ == "__main__":
    main()
