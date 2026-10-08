#!/usr/bin/env python3
"""Calculate simplified SymPy complexity for formulas stored in an Excel column."""

from __future__ import annotations

import argparse
import keyword
import re
from copy import copy
from pathlib import Path
from typing import Any

import sympy as sp
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries


# These defaults can be changed here or overridden with command-line options.
DEFAULT_FORMULA_COLUMN = "GenerationFormula"
DEFAULT_COMPLEXITY_COLUMN = "GenerationFormula_complexity"
DEFAULT_HEADER_ROW = 1


SYMPY_FUNCTIONS: dict[str, Any] = {
    "sqrt": sp.sqrt,
    "exp": sp.exp,
    "log": sp.log,
    "ln": sp.log,
    "sin": sp.sin,
    "cos": sp.cos,
    "tan": sp.tan,
    "asin": sp.asin,
    "acos": sp.acos,
    "atan": sp.atan,
    "asinh": sp.asinh,
    "acosh": sp.acosh,
    "atanh": sp.atanh,
    "sinh": sp.sinh,
    "cosh": sp.cosh,
    "tanh": sp.tanh,
    "Abs": sp.Abs,
    "abs": sp.Abs,
    "Max": sp.Max,
    "Min": sp.Min,
}


IDENTIFIER_PATTERN = re.compile(r"\b[A-Za-z_]\w*\b")


def normalize_formula(formula: str) -> str:
    """Normalize common NumPy/Python spellings used in formula text."""
    text = formula.strip()
    if text.startswith("'"):
        text = text[1:].lstrip()
    if text.startswith("="):
        text = text[1:].lstrip()
    text = re.sub(r"\b(?:numpy|np)\.", "", text)
    text = text.replace("^", "**")
    return text


def build_symbol_table(formula: str) -> tuple[str, dict[str, Any]]:
    """Register physical variable names explicitly to avoid SymPy collisions."""
    normalized = normalize_formula(formula)
    local_dict: dict[str, Any] = dict(SYMPY_FUNCTIONS)
    local_dict["pi"] = sp.pi

    for name in sorted(set(IDENTIFIER_PATTERN.findall(normalized)), key=len, reverse=True):
        if name in local_dict:
            continue

        if keyword.iskeyword(name):
            # For example, physical decay constant "lambda" is a Python
            # keyword, so parse through a safe alias while retaining its name.
            safe_name = f"_symbol_{name}"
            normalized = re.sub(rf"\b{re.escape(name)}\b", safe_name, normalized)
            local_dict[safe_name] = sp.Symbol(name, real=True)
        else:
            # Names such as gamma, S, N and I otherwise collide with SymPy
            # built-ins; they are variables in this scientific spreadsheet.
            local_dict[name] = sp.Symbol(name, real=True)

    return normalized, local_dict


def calculate_complexity(formula: str) -> tuple[int, str]:
    """Return count_ops complexity and the simplified expression text."""
    normalized, local_dict = build_symbol_table(formula)
    expression = sp.sympify(normalized, locals=local_dict)
    simplified = sp.simplify(expression)
    return int(sp.count_ops(simplified, visual=False)), str(simplified)


def find_column(worksheet: Any, header_row: int, column_name: str) -> int | None:
    for cell in worksheet[header_row]:
        if cell.value is not None and str(cell.value).strip() == column_name:
            return cell.column
    return None


def show_progress(completed: int, total: int, message: str) -> None:
    width = 30
    ratio = completed / total if total else 1.0
    filled = min(width, int(width * ratio))
    bar = "█" * filled + "-" * (width - filled)
    status = message.replace("\n", " ")
    if len(status) > 45:
        status = "..." + status[-42:]
    print(
        f"\r[{bar}] {completed:>{len(str(max(total, 1)))}}/{total} "
        f"({ratio * 100:6.2f}%)  {status:<45}",
        end="",
        flush=True,
    )


def copy_cell_style(source: Any, target: Any) -> None:
    """Copy formatting without changing the source workbook's existing cells."""
    if source.has_style:
        target._style = copy(source._style)
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def expand_tables_for_appended_column(
    worksheet: Any,
    header_row: int,
    previous_max_column: int,
    new_column: int,
) -> None:
    """Extend a table by one adjacent column when its range covers the data."""
    if new_column != previous_max_column + 1:
        return

    for table in worksheet.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        if min_row == header_row and max_col == previous_max_column:
            table.ref = (
                f"{get_column_letter(min_col)}{min_row}:"
                f"{get_column_letter(new_column)}{max_row}"
            )


def process_sheet(
    worksheet: Any,
    formula_column_name: str,
    complexity_column_name: str,
    header_row: int,
) -> tuple[int, list[tuple[int, str]]]:
    formula_column = find_column(worksheet, header_row, formula_column_name)
    if formula_column is None:
        raise KeyError(
            f"工作表 {worksheet.title!r} 中没有列 {formula_column_name!r}"
        )

    original_max_column = worksheet.max_column
    complexity_column = find_column(worksheet, header_row, complexity_column_name)
    if complexity_column is None:
        complexity_column = original_max_column + 1
        header_cell = worksheet.cell(header_row, complexity_column)
        header_cell.value = complexity_column_name
        copy_cell_style(worksheet.cell(header_row, formula_column), header_cell)
        worksheet.column_dimensions[get_column_letter(complexity_column)].width = max(
            18, len(complexity_column_name) + 2
        )
        expand_tables_for_appended_column(
            worksheet, header_row, original_max_column, complexity_column
        )

    data_rows = list(range(header_row + 1, worksheet.max_row + 1))
    errors: list[tuple[int, str]] = []
    calculated = 0

    print(f"处理工作表：{worksheet.title}")
    show_progress(0, len(data_rows), "准备开始")
    for position, row_number in enumerate(data_rows, start=1):
        formula_cell = worksheet.cell(row_number, formula_column)
        result_cell = worksheet.cell(row_number, complexity_column)
        copy_cell_style(formula_cell, result_cell)
        result_cell.number_format = "0"
        formula = formula_cell.value

        if formula is None or not str(formula).strip():
            result_cell.value = None
        else:
            try:
                complexity, _ = calculate_complexity(str(formula))
                result_cell.value = complexity
                calculated += 1
            except Exception as exc:
                result_cell.value = None
                errors.append((row_number, f"{type(exc).__name__}: {exc}"))

        show_progress(position, len(data_rows), f"第 {row_number} 行")

    print()
    return calculated, errors


def default_output_path(input_path: Path) -> Path:
    return input_path.with_name(f"{input_path.stem}_with_complexity{input_path.suffix}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="计算 Excel 中 GenerationFormula 的 SymPy 化简复杂度。"
    )
    parser.add_argument("input", type=Path, help="输入 .xlsx 或 .xlsm 文件")
    parser.add_argument("-o", "--output", type=Path, help="输出文件路径")
    parser.add_argument(
        "--formula-column",
        default=DEFAULT_FORMULA_COLUMN,
        help=f"公式列名（默认：{DEFAULT_FORMULA_COLUMN}）",
    )
    parser.add_argument(
        "--complexity-column",
        default=DEFAULT_COMPLEXITY_COLUMN,
        help=f"复杂度列名（默认：{DEFAULT_COMPLEXITY_COLUMN}）",
    )
    parser.add_argument(
        "--sheet",
        help="只处理指定工作表；省略时处理所有包含公式列的工作表",
    )
    parser.add_argument(
        "--header-row",
        type=int,
        default=DEFAULT_HEADER_ROW,
        help=f"表头所在行（默认：{DEFAULT_HEADER_ROW}）",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    input_path = args.input.expanduser().resolve()
    output_path = (
        args.output.expanduser().resolve()
        if args.output is not None
        else default_output_path(input_path)
    )

    if not input_path.is_file():
        raise FileNotFoundError(f"输入文件不存在：{input_path}")
    if input_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("仅支持 .xlsx 和 .xlsm 文件。")
    if args.header_row <= 0:
        raise ValueError("header-row 必须大于 0。")

    keep_vba = input_path.suffix.lower() == ".xlsm"
    workbook = load_workbook(input_path, data_only=False, keep_vba=keep_vba)

    if args.sheet:
        if args.sheet not in workbook.sheetnames:
            raise KeyError(
                f"找不到工作表 {args.sheet!r}；可用工作表：{workbook.sheetnames}"
            )
        worksheets = [workbook[args.sheet]]
    else:
        worksheets = [
            sheet
            for sheet in workbook.worksheets
            if find_column(sheet, args.header_row, args.formula_column) is not None
        ]
        if not worksheets:
            raise KeyError(
                f"所有工作表中都没有找到列 {args.formula_column!r}。"
            )

    total_calculated = 0
    all_errors: list[tuple[str, int, str]] = []
    for worksheet in worksheets:
        calculated, errors = process_sheet(
            worksheet,
            formula_column_name=args.formula_column,
            complexity_column_name=args.complexity_column,
            header_row=args.header_row,
        )
        total_calculated += calculated
        all_errors.extend((worksheet.title, row, message) for row, message in errors)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    print(f"完成：计算 {total_calculated} 个公式的复杂度。")
    if all_errors:
        print(f"有 {len(all_errors)} 行无法解析：")
        for sheet_name, row, message in all_errors:
            print(f"  - {sheet_name}!第 {row} 行：{message}")
    print(f"输出文件：{output_path}")


if __name__ == "__main__":
    main()
