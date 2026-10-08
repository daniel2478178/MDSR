#!/usr/bin/env python3
"""从所有 Pxx/checkpoint.pkl 汇总 database._top_programs 到一个 Excel。

默认使用受限反序列化器读取 checkpoint，不要求当前环境安装 llmsr。
pickle 文件可能包含可执行载荷，因此请只处理自己生成或确认可信的文件。
"""

from __future__ import annotations

import argparse
import ast
import copy
import json
import pickle
import re
from pathlib import Path
from typing import Any, Iterable
from paths import LOGS_ROOT

try:
    import numpy as np
    import openpyxl  # noqa: F401  # pandas 写入 xlsx 时需要
    import pandas as pd
except ModuleNotFoundError as exc:
    raise SystemExit(
        "缺少运行依赖。请先执行：pip install numpy pandas openpyxl"
    ) from exc


# ======================== 常用参数：通常只需修改这里 ========================
PROBLEMS_ROOT = LOGS_ROOT
CHECKPOINT_FILENAME = "checkpoint.pkl"
TOP_N_PER_PROBLEM = 1
SORT_FIELD = "score"
DESCENDING = True
OUTPUT_FILENAME = "all_problems_checkpoint_top1.xlsx"
# ===========================================================================


class _CheckpointPlaceholder:
    """承接 checkpoint 中 llmsr 自定义对象的属性，不运行项目代码。"""


_PLACEHOLDER_CLASSES: dict[tuple[str, str], type] = {}


class RestrictedCheckpointUnpickler(pickle.Unpickler):
    """只允许当前 LLM-SR checkpoint 所需的类和 NumPy构造器。"""

    def find_class(self, module: str, name: str) -> Any:
        if module.startswith("llmsr."):
            key = (module, name)
            if key not in _PLACEHOLDER_CLASSES:
                _PLACEHOLDER_CLASSES[key] = type(
                    name,
                    (_CheckpointPlaceholder,),
                    {"__module__": __name__},
                )
            return _PLACEHOLDER_CLASSES[key]

        numpy_aliases = {
            ("numpy.core.multiarray", "scalar"): np._core.multiarray.scalar,
            ("numpy._core.multiarray", "scalar"): np._core.multiarray.scalar,
            ("numpy", "dtype"): np.dtype,
            ("numpy.core.numeric", "_frombuffer"): np._core.numeric._frombuffer,
            ("numpy._core.numeric", "_frombuffer"): np._core.numeric._frombuffer,
        }
        allowed = numpy_aliases.get((module, name))
        if allowed is not None:
            return allowed
        raise pickle.UnpicklingError(f"checkpoint 包含未允许的对象：{module}.{name}")


def load_checkpoint(path: Path) -> dict[str, Any]:
    with path.open("rb") as file:
        value = RestrictedCheckpointUnpickler(file).load()
    if not isinstance(value, dict):
        raise TypeError("checkpoint 根对象不是字典")
    return value


class FormulaTransformer(ast.NodeTransformer):
    """展开局部变量，并把 params[0]、np.exp 等改成可读形式。"""

    def __init__(self, env: dict[str, ast.AST]) -> None:
        self.env = env

    def visit_Name(self, node: ast.Name) -> ast.AST:
        replacement = self.env.get(node.id)
        return copy.deepcopy(replacement) if replacement is not None else node

    def visit_Subscript(self, node: ast.Subscript) -> ast.AST:
        node = self.generic_visit(node)
        if (
            isinstance(node.value, ast.Name)
            and node.value.id == "params"
            and isinstance(node.slice, ast.Constant)
            and isinstance(node.slice.value, int)
        ):
            return ast.copy_location(ast.Name(id=f"p{node.slice.value}", ctx=ast.Load()), node)
        return node

    def visit_Attribute(self, node: ast.Attribute) -> ast.AST:
        node = self.generic_visit(node)
        if isinstance(node.value, ast.Name) and node.value.id in {"np", "numpy"}:
            return ast.copy_location(ast.Name(id=node.attr, ctx=ast.Load()), node)
        return node


def extract_return_expression(function_source: str) -> ast.AST:
    tree = ast.parse(function_source)
    functions = [node for node in tree.body if isinstance(node, ast.FunctionDef)]
    if not functions:
        raise ValueError("没有找到函数定义")

    env: dict[str, ast.AST] = {}
    for statement in functions[0].body:
        transformer = FormulaTransformer(env)
        if isinstance(statement, ast.Assign):
            value = transformer.visit(copy.deepcopy(statement.value))
            for target in statement.targets:
                if not isinstance(target, ast.Name):
                    raise ValueError("仅支持给普通变量赋值")
                env[target.id] = value
        elif isinstance(statement, ast.AnnAssign) and statement.value is not None:
            if not isinstance(statement.target, ast.Name):
                raise ValueError("仅支持给普通变量赋值")
            env[statement.target.id] = transformer.visit(copy.deepcopy(statement.value))
        elif isinstance(statement, ast.Return):
            if statement.value is None:
                raise ValueError("return 没有返回表达式")
            return FormulaTransformer(env).visit(copy.deepcopy(statement.value))
        elif isinstance(statement, (ast.Expr, ast.Pass)):
            continue
        else:
            raise ValueError(f"不支持的函数体语句：{type(statement).__name__}")
    raise ValueError("函数中没有找到 return 语句")


def full_function_source(program: Any, fallback_body: str = "") -> str:
    body = str(getattr(program, "body", fallback_body) or fallback_body).rstrip()
    if re.search(r"^\s*(?:async\s+)?def\s+\w+\s*\(", body):
        return body.strip()

    name = str(getattr(program, "name", "equation") or "equation")
    args = str(getattr(program, "args", "") or "")
    return_type = str(getattr(program, "return_type", "") or "")
    suffix = f" -> {return_type}" if return_type else ""
    if not body.strip():
        body = "    pass"
    elif not all(line.startswith((" ", "\t")) or not line for line in body.splitlines()):
        body = "\n".join(f"    {line}" if line else "" for line in body.splitlines())
    return f"def {name}({args}){suffix}:\n{body}"


def function_to_formula(function_source: str) -> str:
    expression = ast.fix_missing_locations(extract_return_expression(function_source))
    return ast.unparse(expression)


def natural_key(value: str | Path) -> list[Any]:
    text = value.name if isinstance(value, Path) else str(value)
    return [
        int(part) if part.isdigit() else part.lower()
        for part in re.split(r"(\d+)", text)
    ]


def jsonable(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, dict):
        return {str(key): jsonable(child) for key, child in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [jsonable(child) for child in value]
    if hasattr(value, "__dict__"):
        return {str(key): jsonable(child) for key, child in vars(value).items()}
    return str(value)


def excel_value(value: Any) -> Any:
    converted = jsonable(value)
    if converted is None or isinstance(converted, (str, int, float, bool)):
        return converted
    return json.dumps(converted, ensure_ascii=False, sort_keys=True)


def flatten_mapping(
    mapping: dict[Any, Any], prefix: str = ""
) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in mapping.items():
        column = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(value, dict):
            result.update(flatten_mapping(value, column))
        else:
            result[column] = excel_value(value)
    return result


def iter_island_programs(database: Any) -> Iterable[tuple[int, Any]]:
    for island_index, island in enumerate(getattr(database, "_islands", []) or []):
        clusters = getattr(island, "_clusters", {}) or {}
        cluster_values = clusters.values() if isinstance(clusters, dict) else clusters
        for cluster in cluster_values:
            for program in getattr(cluster, "_programs", []) or []:
                yield island_index, program


def build_program_lookup(database: Any) -> dict[str, list[tuple[int, Any]]]:
    lookup: dict[str, list[tuple[int, Any]]] = {}
    for island_index, program in iter_island_programs(database):
        body = str(getattr(program, "body", "") or "")
        lookup.setdefault(body, []).append((island_index, program))
    return lookup


def safe_numeric(value: Any) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if np.isfinite(number) else None


def find_problem_dirs(problems_root: Path, checkpoint_name: str) -> list[Path]:
    return sorted(
        [
            path
            for path in problems_root.iterdir()
            if path.is_dir()
            and re.fullmatch(r"P\d+", path.name, flags=re.I)
            and (path / checkpoint_name).is_file()
        ],
        key=natural_key,
    )


def collect_problem(
    problem_dir: Path,
    checkpoint_name: str,
    top_n: int,
    sort_field: str,
    descending: bool,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    checkpoint_path = problem_dir / checkpoint_name
    errors: list[dict[str, Any]] = []
    checkpoint = load_checkpoint(checkpoint_path)
    database = checkpoint.get("database")
    if database is None:
        raise KeyError("checkpoint 中没有 database")

    top_programs = getattr(database, "_top_programs", None)
    if not isinstance(top_programs, dict):
        raise TypeError("database._top_programs 不存在或不是字典")

    entries: list[tuple[Any, dict[str, Any]]] = []
    for key, value in top_programs.items():
        if isinstance(value, dict):
            entries.append((key, value))
        else:
            errors.append(
                {
                    "problem_id": problem_dir.name,
                    "stage": "读取 _top_programs",
                    "error": f"跳过非字典记录：{type(value).__name__}",
                }
            )

    def rank_value(item: tuple[Any, dict[str, Any]]) -> tuple[bool, float]:
        value = item[1].get(sort_field)
        if value is None and isinstance(item[1].get("stats"), dict):
            value = item[1]["stats"].get(sort_field)
        number = safe_numeric(value)
        if number is None:
            return (True, 0.0)
        return (False, -number if descending else number)

    entries.sort(key=rank_value)
    entries = entries[:top_n]
    program_lookup = build_program_lookup(database)
    rows: list[dict[str, Any]] = []

    for rank, (top_key, entry) in enumerate(entries, start=1):
        program = entry.get("program")
        if program is None:
            errors.append(
                {
                    "problem_id": problem_dir.name,
                    "top_rank": rank,
                    "stage": "读取程序",
                    "error": "记录中没有 program",
                }
            )
            continue

        body = str(getattr(program, "body", top_key) or top_key)
        function = full_function_source(program, body)
        matches = program_lookup.get(body, [])
        matched_orders = sorted(
            {
                int(order)
                for _, matched in matches
                if (order := safe_numeric(getattr(matched, "global_sample_nums", None)))
                is not None
            }
        )
        island_ids = sorted({island_index for island_index, _ in matches})

        stats = entry.get("stats") if isinstance(entry.get("stats"), dict) else {}
        scores_per_test = (
            entry.get("scores_per_test")
            if isinstance(entry.get("scores_per_test"), dict)
            else {}
        )
        row: dict[str, Any] = {
            "problem_id": problem_dir.name,
            "top_rank": rank,
            "checkpoint_file": str(checkpoint_path),
            "checkpoint_global_sample_nums": excel_value(
                checkpoint.get("global_sample_nums")
            ),
            "database_top_k": excel_value(getattr(database, "_top_k", None)),
            "sample_order": matched_orders[0] if len(matched_orders) == 1 else None,
            "matched_sample_orders": json.dumps(matched_orders, ensure_ascii=False),
            "matched_island_ids": json.dumps(island_ids, ensure_ascii=False),
            "function": function,
            "formula": None,
            "program_body": body,
            "top_program_key": excel_value(top_key),
            "simplified_formula": excel_value(entry.get("simplified_formula")),
            "score": excel_value(entry.get("score")),
            "complexity": excel_value(entry.get("complexity")),
            "mean_nmse": excel_value(stats.get("mean_nmse")),
            "min_nmse": excel_value(stats.get("min_nmse")),
            "max_nmse": excel_value(stats.get("max_nmse")),
            "median_nmse": excel_value(stats.get("median_nmse")),
            "std_nmse": excel_value(stats.get("std_nmse")),
            "worst2_mean_nmse": excel_value(stats.get("worst2_mean_nmse")),
            "scores_per_test_json": excel_value(scores_per_test),
        }

        try:
            row["formula"] = function_to_formula(function)
        except (SyntaxError, TypeError, ValueError, NotImplementedError) as exc:
            errors.append(
                {
                    "problem_id": problem_dir.name,
                    "top_rank": rank,
                    "stage": "function 转公式",
                    "error": str(exc),
                }
            )

        # 保存记录、stats、每个数据集得分以及 program 对象中的全部字段。
        row.update(flatten_mapping(stats, "stats"))
        row.update(flatten_mapping(scores_per_test, "scores_per_test"))
        row.update(flatten_mapping(vars(program), "program"))
        extras = {
            key: value
            for key, value in entry.items()
            if key not in {"program", "stats", "scores_per_test"}
        }
        row.update(flatten_mapping(extras, "entry"))
        rows.append(row)

    return rows, errors


PREFERRED_COLUMNS = [
    "problem_id",
    "top_rank",
    "sample_order",
    "function",
    "formula",
    "simplified_formula",
    "score",
    "complexity",
    "mean_nmse",
    "min_nmse",
    "max_nmse",
    "median_nmse",
    "std_nmse",
    "worst2_mean_nmse",
    "scores_per_test_json",
    "matched_sample_orders",
    "matched_island_ids",
    "checkpoint_global_sample_nums",
    "database_top_k",
    "checkpoint_file",
    "program_body",
    "top_program_key",
]


def write_excel(
    rows: list[dict[str, Any]],
    errors: list[dict[str, Any]],
    output_file: Path,
) -> None:
    frame = pd.DataFrame(rows)
    leading = [column for column in PREFERRED_COLUMNS if column in frame.columns]
    remaining = [column for column in frame.columns if column not in leading]
    frame = frame[leading + remaining]

    summary = (
        frame.groupby("problem_id", as_index=False)
        .agg(program_count=("top_rank", "count"))
        .sort_values("problem_id", key=lambda series: series.map(str))
    )

    output_file.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="Top_Programs", index=False)
        summary.to_excel(writer, sheet_name="Summary", index=False)
        if errors:
            pd.DataFrame(errors).to_excel(writer, sheet_name="Errors", index=False)

        from openpyxl.styles import Alignment, Font, PatternFill

        for sheet in writer.sheets.values():
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="4472C4")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for column_cells in sheet.columns:
                letter = column_cells[0].column_letter
                header = str(column_cells[0].value or "")
                if header in {"function", "program_body", "top_program_key", "program.docstring"}:
                    width = 70
                elif header in {"formula", "simplified_formula", "scores_per_test_json"}:
                    width = 50
                else:
                    width = min(
                        max(len(str(cell.value or "")) for cell in column_cells) + 2,
                        28,
                    )
                sheet.column_dimensions[letter].width = max(width, 12)
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="从所有 Pxx/checkpoint.pkl 的 _top_programs 汇总最佳公式。"
    )
    parser.add_argument(
        "--problems-root",
        type=Path,
        default=PROBLEMS_ROOT,
        help=f"包含 P01、P02……的目录（默认：{PROBLEMS_ROOT}）",
    )
    parser.add_argument(
        "--checkpoint-name",
        default=CHECKPOINT_FILENAME,
        help=f"checkpoint 文件名（默认：{CHECKPOINT_FILENAME}）",
    )
    parser.add_argument(
        "-n",
        "--top-n",
        type=int,
        default=TOP_N_PER_PROBLEM,
        help=f"每个问题读取的程序数（默认：{TOP_N_PER_PROBLEM}）",
    )
    parser.add_argument(
        "--sort-field",
        default=SORT_FIELD,
        help=f"排序字段，可用 score 或 stats 中的字段（默认：{SORT_FIELD}）",
    )
    direction = parser.add_mutually_exclusive_group()
    direction.add_argument("--descending", dest="descending", action="store_true")
    direction.add_argument("--ascending", dest="descending", action="store_false")
    parser.set_defaults(descending=DESCENDING)
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=None,
        help="输出 Excel；默认保存在 problems-root 中",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.top_n <= 0:
        raise ValueError("top_n 必须大于 0")

    problems_root = args.problems_root.expanduser().resolve()
    if not problems_root.is_dir():
        raise NotADirectoryError(f"问题根目录不存在：{problems_root}")

    problem_dirs = find_problem_dirs(problems_root, args.checkpoint_name)
    if not problem_dirs:
        raise FileNotFoundError(
            f"{problems_root} 下没有找到 Pxx/{args.checkpoint_name}"
        )

    output_file = (
        args.output.expanduser().resolve()
        if args.output is not None
        else problems_root / OUTPUT_FILENAME
    )

    all_rows: list[dict[str, Any]] = []
    all_errors: list[dict[str, Any]] = []
    for index, problem_dir in enumerate(problem_dirs, start=1):
        print(f"[{index}/{len(problem_dirs)}] 处理 {problem_dir.name} ...")
        try:
            rows, errors = collect_problem(
                problem_dir,
                args.checkpoint_name,
                args.top_n,
                args.sort_field,
                args.descending,
            )
            all_rows.extend(rows)
            all_errors.extend(errors)
            print(f"    读取 {len(rows)} 个 Top 程序")
        except Exception as exc:  # 单个问题损坏时继续处理其他问题
            all_errors.append(
                {
                    "problem_id": problem_dir.name,
                    "stage": "读取 checkpoint",
                    "error": f"{type(exc).__name__}: {exc}",
                }
            )
            print(f"    失败：{type(exc).__name__}: {exc}")

    if not all_rows:
        raise ValueError("没有从任何 checkpoint 读取到 _top_programs")

    write_excel(all_rows, all_errors, output_file)
    print(
        f"完成：{len(problem_dirs)} 个问题，共 {len(all_rows)} 个程序。\n"
        f"输出文件：{output_file}"
    )


if __name__ == "__main__":
    main()
