#!/usr/bin/env python3
"""汇总所有 Pxx 问题中指标最优的样本，并将 function 转为公式。"""

from __future__ import annotations

import argparse
import ast
import json
import re
from pathlib import Path
from typing import Any
from paths import LOGS_ROOT

try:
    import openpyxl  # noqa: F401  # pandas 写入 xlsx 时需要
    import pandas as pd
    import sympy as sp
except ModuleNotFoundError as exc:
    raise SystemExit(
        "缺少运行依赖。请先执行：pip install pandas openpyxl sympy"
    ) from exc


# ======================== 常用参数：通常只需修改这里 ========================
P01_SAMPLES_DIR = LOGS_ROOT / "P01" / "samples"
TOP_N_PER_PROBLEM = 10
RANK_FIELD = "mean_nmse"
SMALLER_IS_BETTER = True
# ===========================================================================

OUTPUT_COLUMNS = [
    "problem_id",
    "sample_order",
    "function",
    "formula",
    "score",
    "mean_nmse",
    "min_nmse",
    "max_nmse",
    "std_nmse",
]


class SymbolicFunctionParser:
    """将受限的 NumPy 风格 Python 函数安全转换为 SymPy 表达式。"""

    FUNCTIONS = {
        "sin": sp.sin,
        "cos": sp.cos,
        "tan": sp.tan,
        "arcsin": sp.asin,
        "asin": sp.asin,
        "arccos": sp.acos,
        "acos": sp.acos,
        "arctan": sp.atan,
        "atan": sp.atan,
        "sinh": sp.sinh,
        "cosh": sp.cosh,
        "tanh": sp.tanh,
        "exp": sp.exp,
        "log": sp.log,
        "sqrt": sp.sqrt,
        "abs": sp.Abs,
        "fabs": sp.Abs,
        "sign": sp.sign,
        "maximum": sp.Max,
        "minimum": sp.Min,
    }

    def __init__(self) -> None:
        self.env: dict[str, sp.Expr] = {}

    @staticmethod
    def dotted_name(node: ast.AST) -> str:
        if isinstance(node, ast.Name):
            return node.id
        if isinstance(node, ast.Attribute):
            left = SymbolicFunctionParser.dotted_name(node.value)
            return f"{left}.{node.attr}" if left else node.attr
        return ""

    def expression(self, node: ast.AST) -> sp.Expr:
        if isinstance(node, ast.Constant):
            if isinstance(node.value, bool):
                return sp.true if node.value else sp.false
            if isinstance(node.value, int):
                return sp.Integer(node.value)
            if isinstance(node.value, float):
                if node.value.is_integer():
                    return sp.Integer(int(node.value))
                return sp.Float(node.value)
            raise ValueError(f"不支持的常量：{node.value!r}")

        if isinstance(node, ast.Name):
            if node.id == "pi":
                return sp.pi
            return self.env.get(node.id, sp.Symbol(node.id, real=True))

        if isinstance(node, ast.Attribute):
            name = self.dotted_name(node)
            if name in {"np.pi", "numpy.pi"}:
                return sp.pi
            if name in {"np.e", "numpy.e"}:
                return sp.E
            raise ValueError(f"不支持的属性：{name}")

        if isinstance(node, ast.BinOp):
            left = self.expression(node.left)
            right = self.expression(node.right)
            operations = {
                ast.Add: lambda: left + right,
                ast.Sub: lambda: left - right,
                ast.Mult: lambda: left * right,
                ast.Div: lambda: left / right,
                ast.Pow: lambda: left**right,
                ast.Mod: lambda: sp.Mod(left, right),
            }
            operation = operations.get(type(node.op))
            if operation is None:
                raise ValueError(f"不支持的二元运算：{type(node.op).__name__}")
            return operation()

        if isinstance(node, ast.UnaryOp):
            value = self.expression(node.operand)
            if isinstance(node.op, ast.USub):
                return -value
            if isinstance(node.op, ast.UAdd):
                return value
            if isinstance(node.op, ast.Not):
                return sp.Not(value)
            raise ValueError(f"不支持的一元运算：{type(node.op).__name__}")

        if isinstance(node, ast.Subscript):
            base = self.dotted_name(node.value)
            if isinstance(node.slice, ast.Constant) and isinstance(node.slice.value, int):
                index: int | sp.Expr = node.slice.value
            else:
                index = self.expression(node.slice)
            if base == "params" and isinstance(index, int):
                return sp.Symbol(f"p{index}", real=True)
            return sp.Symbol(f"{base}[{index}]", real=True)

        if isinstance(node, ast.Call):
            full_name = self.dotted_name(node.func)
            short_name = full_name.rsplit(".", 1)[-1]
            args = [self.expression(arg) for arg in node.args]
            if short_name in self.FUNCTIONS:
                return self.FUNCTIONS[short_name](*args)
            if short_name == "square" and len(args) == 1:
                return args[0] ** 2
            if short_name in {"pow", "power"} and len(args) == 2:
                return args[0] ** args[1]
            if short_name == "where" and len(args) == 3:
                return sp.Piecewise((args[1], args[0]), (args[2], True))
            if short_name == "clip" and len(args) == 3:
                return sp.Min(sp.Max(args[0], args[1]), args[2])
            raise ValueError(f"不支持的函数调用：{full_name}")

        if isinstance(node, ast.Compare) and len(node.ops) == len(node.comparators) == 1:
            left = self.expression(node.left)
            right = self.expression(node.comparators[0])
            relations = {
                ast.Lt: sp.Lt,
                ast.LtE: sp.Le,
                ast.Gt: sp.Gt,
                ast.GtE: sp.Ge,
                ast.Eq: sp.Eq,
                ast.NotEq: sp.Ne,
            }
            relation = relations.get(type(node.ops[0]))
            if relation is None:
                raise ValueError("不支持的比较运算")
            return relation(left, right)

        raise ValueError(f"不支持的表达式节点：{type(node).__name__}")

    def assign(self, target: ast.AST, value_node: ast.AST) -> None:
        if isinstance(target, (ast.Tuple, ast.List)):
            if not isinstance(value_node, (ast.Tuple, ast.List)):
                raise ValueError("多变量赋值的左右数量无法对应")
            if len(target.elts) != len(value_node.elts):
                raise ValueError("多变量赋值的左右数量不一致")
            for child_target, child_value in zip(target.elts, value_node.elts):
                self.assign(child_target, child_value)
            return
        if not isinstance(target, ast.Name):
            raise ValueError(f"不支持的赋值目标：{type(target).__name__}")
        self.env[target.id] = self.expression(value_node)

    def parse(self, source: str) -> sp.Expr:
        source = clean_function_source(source)
        tree = ast.parse(source)
        functions = [
            node
            for node in tree.body
            if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))
        ]
        if not functions:
            raise ValueError("没有找到函数定义")

        # 按函数体的真实顺序处理，因此 return 前面的中间表达式会先被展开。
        for statement in functions[0].body:
            if isinstance(statement, ast.Assign):
                for target in statement.targets:
                    self.assign(target, statement.value)
            elif isinstance(statement, ast.AnnAssign) and statement.value is not None:
                self.assign(statement.target, statement.value)
            elif isinstance(statement, ast.AugAssign) and isinstance(statement.target, ast.Name):
                old = self.env.get(
                    statement.target.id, sp.Symbol(statement.target.id, real=True)
                )
                right = self.expression(statement.value)
                if isinstance(statement.op, ast.Add):
                    self.env[statement.target.id] = old + right
                elif isinstance(statement.op, ast.Sub):
                    self.env[statement.target.id] = old - right
                elif isinstance(statement.op, ast.Mult):
                    self.env[statement.target.id] = old * right
                elif isinstance(statement.op, ast.Div):
                    self.env[statement.target.id] = old / right
                else:
                    raise ValueError(
                        f"不支持的复合赋值：{type(statement.op).__name__}"
                    )
            elif isinstance(statement, ast.Return):
                if statement.value is None:
                    raise ValueError("return 没有返回表达式")
                return sp.simplify(self.expression(statement.value))
            elif isinstance(statement, (ast.Expr, ast.Pass)):
                continue
            else:
                raise ValueError(f"不支持的函数体语句：{type(statement).__name__}")

        raise ValueError("函数中没有找到 return 语句")


def clean_function_source(source: str) -> str:
    """去除模型输出中可能存在的 Markdown 代码围栏。"""
    text = source.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:python)?\s*", "", text, count=1, flags=re.I)
        text = re.sub(r"\s*```$", "", text, count=1)
    return text.strip()


def function_to_formula(source: str) -> str:
    expression = SymbolicFunctionParser().parse(source)
    return str(expression)


def natural_key(value: str | Path) -> list[Any]:
    text = value.name if isinstance(value, Path) else str(value)
    return [
        int(part) if part.isdigit() else part.lower()
        for part in re.split(r"(\d+)", text)
    ]


def flatten(value: Any, prefix: str = "") -> dict[str, Any]:
    if isinstance(value, dict):
        result: dict[str, Any] = {}
        for key, child in value.items():
            column = f"{prefix}.{key}" if prefix else str(key)
            result.update(flatten(child, column))
        return result
    if isinstance(value, list):
        return {prefix or "value": json.dumps(value, ensure_ascii=False)}
    return {prefix or "value": value}


def pick_field(row: dict[str, Any], field: str) -> Any:
    """优先取顶层字段；也兼容 flatten 后的 stats.mean_nmse 等字段。"""
    if field in row:
        return row[field]
    matches = [value for key, value in row.items() if key.endswith(f".{field}")]
    return matches[0] if len(matches) == 1 else None


def read_json_records(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig") as file:
        data = json.load(file)
    items = data if isinstance(data, list) else [data]
    records: list[dict[str, Any]] = []
    for item in items:
        if not isinstance(item, dict):
            raise ValueError("JSON 根记录不是对象")
        records.append(flatten(item))
    return records


def find_problem_dirs(problems_root: Path) -> list[Path]:
    """查找 P01、P02……等包含 samples 文件夹的问题目录。"""
    problem_dirs = [
        path
        for path in problems_root.iterdir()
        if path.is_dir()
        and re.fullmatch(r"P\d+", path.name, flags=re.I)
        and (path / "samples").is_dir()
    ]
    return sorted(problem_dirs, key=natural_key)


def collect_problem(
    problem_dir: Path,
    top_n: int,
    rank_field: str,
    ascending: bool,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    sample_dir = problem_dir / "samples"
    json_files = sorted(sample_dir.glob("*.json"), key=natural_key)
    candidates: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []

    for json_path in json_files:
        try:
            for source_index, row in enumerate(read_json_records(json_path)):
                record = {
                    "problem_id": problem_dir.name,
                    "source_file": json_path.name,
                    "source_index": source_index,
                }
                for field in set(OUTPUT_COLUMNS + [rank_field]):
                    if field not in {"problem_id", "formula"}:
                        record[field] = pick_field(row, field)
                candidates.append(record)
        except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as exc:
            errors.append(
                {
                    "problem_id": problem_dir.name,
                    "source_file": json_path.name,
                    "stage": "读取 JSON",
                    "error": str(exc),
                }
            )

    if not candidates:
        return [], errors

    frame = pd.DataFrame(candidates)
    numeric_rank = pd.to_numeric(frame[rank_field], errors="coerce")
    invalid_count = int(numeric_rank.isna().sum())
    if invalid_count:
        errors.append(
            {
                "problem_id": problem_dir.name,
                "source_file": "",
                "stage": "筛选",
                "error": f"{invalid_count} 条记录的 {rank_field} 不是有效数值，已跳过",
            }
        )

    frame = (
        frame.loc[numeric_rank.notna()]
        .assign(_rank_numeric=numeric_rank.loc[numeric_rank.notna()])
        .sort_values("_rank_numeric", ascending=ascending, kind="stable")
        .head(top_n)
        .drop(columns="_rank_numeric")
        .reset_index(drop=True)
    )

    selected: list[dict[str, Any]] = []
    for row in frame.to_dict(orient="records"):
        function_source = row.get("function")
        if isinstance(function_source, str) and function_source.strip():
            try:
                row["formula"] = function_to_formula(function_source)
            except (SyntaxError, TypeError, ValueError, NotImplementedError) as exc:
                row["formula"] = None
                errors.append(
                    {
                        "problem_id": problem_dir.name,
                        "source_file": row.get("source_file", ""),
                        "sample_order": row.get("sample_order"),
                        "stage": "function 转公式",
                        "error": str(exc),
                    }
                )
        else:
            row["formula"] = None
            errors.append(
                {
                    "problem_id": problem_dir.name,
                    "source_file": row.get("source_file", ""),
                    "sample_order": row.get("sample_order"),
                    "stage": "function 转公式",
                    "error": "function 字段为空或不是字符串",
                }
            )
        selected.append(row)

    return selected, errors


def write_excel(
    rows: list[dict[str, Any]], errors: list[dict[str, Any]], output_file: Path
) -> None:
    frame = pd.DataFrame(rows)
    for column in OUTPUT_COLUMNS:
        if column not in frame.columns:
            frame[column] = None
    frame = frame[OUTPUT_COLUMNS]

    output_file.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="Top_Samples", index=False)
        sheet = writer.sheets["Top_Samples"]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        from openpyxl.styles import Alignment, Font, PatternFill

        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        widths = {
            "A": 12,
            "B": 16,
            "C": 72,
            "D": 55,
            "E": 16,
            "F": 18,
            "G": 18,
            "H": 18,
            "I": 18,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        if errors:
            error_frame = pd.DataFrame(errors)
            error_frame.to_excel(writer, sheet_name="Errors", index=False)
            error_sheet = writer.sheets["Errors"]
            error_sheet.freeze_panes = "A2"
            error_sheet.auto_filter.ref = error_sheet.dimensions
            for cell in error_sheet[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="C00000")
            for cells in error_sheet.columns:
                letter = cells[0].column_letter
                width = min(max(len(str(cell.value or "")) for cell in cells) + 2, 90)
                error_sheet.column_dimensions[letter].width = max(width, 14)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="按问题汇总指标最优的 JSON 样本，并将 function 转为公式。"
    )
    parser.add_argument(
        "--p01-samples-dir",
        type=Path,
        default=P01_SAMPLES_DIR,
        help=f"P01 的 samples 文件夹（默认：{P01_SAMPLES_DIR}）",
    )
    parser.add_argument(
        "-n",
        "--top-n",
        type=int,
        default=TOP_N_PER_PROBLEM,
        help=f"每个问题保留的样本数（默认：{TOP_N_PER_PROBLEM}）",
    )
    parser.add_argument(
        "--rank-field",
        default=RANK_FIELD,
        help=f"排名字段（默认：{RANK_FIELD}）",
    )
    direction_group = parser.add_mutually_exclusive_group()
    direction_group.add_argument(
        "--descending",
        dest="descending",
        action="store_true",
        help="按排名字段从高到低选取",
    )
    direction_group.add_argument(
        "--ascending",
        dest="descending",
        action="store_false",
        help="按排名字段从低到高选取",
    )
    parser.set_defaults(descending=not SMALLER_IS_BETTER)
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=None,
        help="输出 Excel 路径；默认保存在 P01 的父目录",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.top_n <= 0:
        raise ValueError("top_n 必须是大于 0 的整数")

    p01_samples_dir = args.p01_samples_dir.expanduser().resolve()
    if not p01_samples_dir.is_dir():
        raise NotADirectoryError(f"P01 samples 文件夹不存在：{p01_samples_dir}")

    # ...\logs\P01\samples -> ...\logs，即 P01 的父目录。
    problems_root = p01_samples_dir.parent.parent
    problem_dirs = find_problem_dirs(problems_root)
    if not problem_dirs:
        raise FileNotFoundError(f"没有找到 Pxx\\samples 文件夹：{problems_root}")

    output_file = args.output
    if output_file is None:
        direction = "highest" if args.descending else "lowest"
        output_file = problems_root / (
            f"all_problems_top{args.top_n}_{direction}_{args.rank_field}.xlsx"
        )
    else:
        output_file = output_file.expanduser().resolve()

    all_rows: list[dict[str, Any]] = []
    all_errors: list[dict[str, Any]] = []
    for index, problem_dir in enumerate(problem_dirs, start=1):
        print(f"[{index}/{len(problem_dirs)}] 处理 {problem_dir.name} ...")
        rows, errors = collect_problem(
            problem_dir=problem_dir,
            top_n=args.top_n,
            rank_field=args.rank_field,
            ascending=not args.descending,
        )
        all_rows.extend(rows)
        all_errors.extend(errors)
        print(f"    入选 {len(rows)} 条")

    if not all_rows:
        raise ValueError("所有问题中都没有找到可用于排名的记录")

    write_excel(all_rows, all_errors, output_file)
    print(
        f"完成：{len(problem_dirs)} 个问题，共写入 {len(all_rows)} 条记录。\n"
        f"输出文件：{output_file}"
    )


if __name__ == "__main__":
    main()
